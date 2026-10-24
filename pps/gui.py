"""Tk desktop UI: keyboard input, Excel import, reports, tables and plots."""
import csv
import json
import traceback
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
import numpy as np
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from openpyxl import load_workbook, Workbook
from .common import numbers, fmt, set_display_digits
from . import interpolation as ip, calculus as ca, differential as de
from .math_view import MathDocument, MathTable, export_html, table_display

ROOT=Path(__file__).resolve().parents[1]

# Fields: name, label, default, help. Every default is an editable example.
def F(key,label,default='',help=''):
    return key,label,str(default),help

TOOLS={
 'Nội suy đa thức':(ip.METHODS,True),
 'Newton':(['Mốc cách đều','Mốc bất kỳ'],True),
 'Spline':(['Tuyến tính','Bậc hai — biên trái','Bậc hai — biên phải','Cubic tự nhiên','Cubic kẹp','Cubic đạo hàm bậc hai'],True),
 'Bình phương tối thiểu':(['Tuyến tính theo cơ sở','Lũy thừa','Hàm mũ'],True),
 'Nội suy ngược':(['Lặp tiến','Lặp lùi'],True),
 'Chọn mốc':(['Mốc cách đều','Mốc nội suy ngược','Chebyshev'],True),
 'Đạo hàm & Horner':(['Đạo hàm từ bảng','Horner'],True),
 'Tích phân':(['Hình thang','Simpson','Trung điểm','Newton-Cotes','Tích phân kép'],False),
 'Phương trình vi phân':(['Euler hiện','Euler ẩn','Hình thang','RK2','RK3','RK4','Adams 2','Adams 3','Adams 4','Adams 5','RK2 tham số','RK3 tham số','RK tùy chỉnh'],False),
 'Bài toán biên':(['Biên tổng quát','Trị riêng'],False),
}


def fields(tool,method):
    if tool=='Newton': return [F('direction','Hướng nội suy','Tiến'),F('target','Điểm cần tính x',1.5),F('order','Cấp đạo hàm (0 = giá trị hàm)',0)]
    if tool=='Nội suy đa thức': return [F('target','Điểm cần tính x',1.5),F('order','Cấp đạo hàm (0 = giá trị hàm)',0)]
    if tool=='Spline':
        v=[F('target','Điểm cần tính x',1.5)]
        if method in ('Bậc hai — biên trái','Cubic kẹp','Cubic đạo hàm bậc hai'): v.append(F('left','Giá trị đạo hàm tại biên trái',0))
        if method in ('Bậc hai — biên phải','Cubic kẹp','Cubic đạo hàm bậc hai'): v.append(F('right','Giá trị đạo hàm tại biên phải',0))
        return v
    if tool=='Bình phương tối thiểu':
        return [] if method=='Lũy thừa' else [F('basis','Các hàm cơ sở (ngăn bằng dấu ;)', '1; x; x^2','Hàm mũ tự thêm hệ số hằng trong số mũ.')]
    if tool=='Nội suy ngược': return [F('target','Giá trị y cần tìm ngược',1.3225),F('tol','Dung sai lặp',1e-10)]
    if tool=='Chọn mốc':
        if method=='Chebyshev': return [F('degree','Bậc đa thức n (sinh n+1 mốc)',4),F('a','Đầu đoạn a',-1),F('b','Cuối đoạn b',1)]
        return [F('target','Giá trị y' if method=='Mốc nội suy ngược' else 'Điểm x',1.5)]+([] if method=='Mốc nội suy ngược' else [F('count','Số mốc cần lấy',3)])
    if tool=='Đạo hàm & Horner':
        if method=='Horner': return [F('coef','Hệ số theo bậc giảm dần','1 0 0','Ví dụ 1 0 0 biểu diễn x².'),F('target','Điểm x',1.5),F('order','Đạo hàm đến cấp',2)]
        return [F('target','Điểm x',1.5),F('degree','Bậc p (dùng p+1 mốc)',2),F('order','Cấp đạo hàm',1)]
    if tool=='Tích phân':
        v=[F('expr','Hàm f(x,y)' if method=='Tích phân kép' else 'Hàm f(x)','x^2+y^2' if method=='Tích phân kép' else 'x^2'),F('a','Cận trái a',0),F('b','Cận phải b',1),F('n','Số khoảng chia n',20)]
        if method=='Tích phân kép':
            return v+[F('c','Cận dưới y',0),F('d','Cận trên y',1),F('rule','Quy tắc','Simpson','Nhập Simpson hoặc Hình thang.'),F('tol','Dung sai Runge',1e-7)]
        if method=='Newton-Cotes': v.append(F('degree','Bậc Newton-Cotes',4,'n phải chia hết cho bậc này.'))
        return v+[F('tol','Dung sai tự chia lưới (để trống: dùng n)','','Chỉ dùng với hàm; tự tăng n đến khi ước lượng Runge đủ nhỏ.'),F('bound','Cận đạo hàm M (không bắt buộc)','','M2 cho hình thang/trung điểm, M4 cho Simpson.'),F('source','Nguồn dữ liệu','Hàm','Nhập Hàm hoặc Bảng; bảng x,y điền bên dưới.')]
    if tool=='Phương trình vi phân':
        v=[F('expr','Vế phải x′; y′; z′ (1–3 hàm)','x','Biến độc lập t; một phương trình x′=x nhập x. Hệ: y; -x'),F('initial','Giá trị đầu x, y, z','1'),F('a','Thời điểm đầu t0',0),F('b','Thời điểm cuối',1),F('h','Bước h',.1),F('tol','Dung sai bước ẩn',1e-10)]
        if method=='RK2 tham số': v += [F('parameter','Tham số biết trước','alpha2','alpha2, beta11, r1 hoặc r2.'),F('value','Giá trị tham số',1)]
        if method=='RK3 tham số': v += [F('alpha2','alpha2',.5),F('alpha3','alpha3',1)]
        if method=='RK tùy chỉnh':
            v += [F('rk_a','Ma trận A (JSON)','[[0,0],[1,0]]'),F('rk_b','Trọng số b','0.5 0.5'),F('rk_c','Vị trí c','0 1','Cho RK2/RK3 hoặc RK hiện nhiều tầng; A tam giác dưới.')]
        return v
    if tool=='Bài toán biên':
        v=[F('p','Hàm p(x)','1'),F('q','Hàm q(x)','0'),F('f','Hàm r(x)' if method=='Trị riêng' else 'Hàm f(x)','1' if method=='Trị riêng' else '2'),F('a','Đầu đoạn a',0),F('b','Cuối đoạn b',1),F('h','Bước h',.05)]
        if method=='Trị riêng': return v+[F('count','Số trị riêng nhỏ nhất',5)]
        return v+[F('left','Biên trái: alpha beta gamma','0 1 0',"alpha*u′(a) + beta*u(a) = gamma"),F('right','Biên phải: alpha beta gamma','0 1 0',"alpha*u′(b) + beta*u(b) = gamma")]
    return []


def dispatch(tool,method,values,xtext,ytext):
    def num(k,default=0):
        v=float(values.get(k,str(default)))
        if not np.isfinite(v): raise ValueError(f'{k} phải hữu hạn.')
        return v
    def integer(k,default=0):
        v=num(k,default)
        if int(v)!=v: raise ValueError(f'{k} phải là số nguyên.')
        return int(v)
    def optional(k): return num(k) if values.get(k,'').strip() else None
    def data(): return numbers(xtext),numbers(ytext)
    if tool=='Newton':
        backward=values.get('direction','Tiến')=='Lùi'
        method=('Newton lùi' if backward else 'Newton tiến') if method=='Mốc cách đều' else ('Newton bất kỳ lùi' if backward else 'Newton bất kỳ')
        return ip.interpolate(*data(),method,num('target'),integer('order'))
    if tool=='Nội suy đa thức': return ip.interpolate(*data(),method,num('target'),integer('order'))
    if tool=='Spline': return ip.spline(*data(),method,num('left'),num('right'),num('target'))
    if tool=='Bình phương tối thiểu': return ip.least_squares(*data(),method,[v.strip() for v in values.get('basis','1;x').split(';') if v.strip()])
    if tool=='Nội suy ngược': return ip.inverse(*data(),num('target'),method=='Lặp lùi',num('tol'))
    if tool=='Chọn mốc':
        if method=='Chebyshev': return ip.chebyshev(integer('degree'),num('a'),num('b'))
        return ip.select_points(*data(),num('target'),integer('count',3),method=='Mốc nội suy ngược')
    if tool=='Đạo hàm & Horner':
        if method=='Horner': return ip.horner(numbers(values['coef']),num('target'),integer('order'))
        return ca.differentiate(*data(),num('target'),integer('degree'),integer('order'))
    if tool=='Tích phân':
        if method=='Tích phân kép': return ca.double_integral(values['expr'],num('a'),num('b'),num('c'),num('d'),integer('n'),values['rule'].strip(),num('tol'))
        source=values['source'].strip().casefold()
        if source not in ('hàm','bảng'): raise ValueError('Nguồn dữ liệu phải là Hàm hoặc Bảng.')
        x,y=data() if source=='bảng' else (None,None)
        return ca.integrate(method,values['expr'],num('a'),num('b'),integer('n'),integer('degree',4),optional('tol'),x,y,optional('bound'))
    if tool=='Phương trình vi phân':
        custom=(json.loads(values['rk_a']),numbers(values['rk_b']),numbers(values['rk_c'])) if method=='RK tùy chỉnh' else None
        if method=='RK2 tham số': custom=de.custom_rk2(values['parameter'].strip(),num('value'));method='RK tùy chỉnh'
        if method=='RK3 tham số': custom=de.custom_rk3(num('alpha2'),num('alpha3'));method='RK tùy chỉnh'
        return de.ode([s.strip() for s in values['expr'].split(';') if s.strip()],numbers(values['initial']),num('a'),num('b'),num('h'),method,num('tol'),custom)
    if tool=='Bài toán biên':
        if method=='Trị riêng': return de.eigen(values['p'],values['q'],values['f'],num('a'),num('b'),num('h'),integer('count'))
        return de.bvp(values['p'],values['q'],values['f'],num('a'),num('b'),num('h'),numbers(values['left']),numbers(values['right']))
    raise ValueError('Chọn chức năng trước khi tính.')


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('PPS • Phương pháp số | Python')
        self.geometry('1380x860'); self.minsize(1050,680)
        self.configure(bg='#f3f5f7')
        self.option_add('*Font',('Segoe UI',10))
        style=ttk.Style(self); style.theme_use('clam')
        style.configure('.',font=('Segoe UI',10),background='#f3f5f7',foreground='#172b3a')
        style.configure('TButton',padding=(12,7))
        style.configure('Accent.TButton',background='#126b64',foreground='white',font=('Segoe UI',11,'bold'))
        style.map('Accent.TButton',background=[('active','#0b554f'),('disabled','#718c89')])
        style.configure('TEntry',fieldbackground='white',padding=5)
        style.configure('TCombobox',fieldbackground='white',padding=5)
        style.configure('TNotebook.Tab',padding=(14,8))
        style.configure('Treeview',background='white',fieldbackground='white',rowheight=28)
        style.configure('Treeview.Heading',font=('Segoe UI',10,'bold'),padding=5)
        self.executor=ThreadPoolExecutor(max_workers=1); self.future=None; self.report=None; self.form={}; self.cache={}
        self.current_tool=None; self.current_method=None
        self.protocol('WM_DELETE_WINDOW',self.close)
        self.build()
        self.bind('<Control-Return>',lambda e:self.run())
        self.bind('<Control-s>',lambda e:self.save_project())
        self.select_tool('Nội suy đa thức')

    def build(self):
        header=tk.Frame(self,bg='#123c43',height=76); header.pack(fill='x'); header.pack_propagate(False)
        tk.Label(header,text='PPS',font=('Segoe UI',24,'bold'),fg='white',bg='#123c43').pack(side='left',padx=(22,16))
        tk.Label(header,text='PHƯƠNG PHÁP SỐ\nNội suy · Tích phân · Phương trình vi phân',justify='left',font=('Segoe UI',10),fg='#e0edef',bg='#123c43').pack(side='left')
        for label,cmd in [('Hướng dẫn',self.help),('Mở bài',self.load_project),('Lưu bài',self.save_project)]:
            ttk.Button(header,text=label,command=cmd).pack(side='right',padx=6,pady=18)
        main=ttk.Panedwindow(self,orient='horizontal'); main.pack(fill='both',expand=True,padx=14,pady=14)
        nav=ttk.Frame(main,width=208); main.add(nav,weight=0)
        ttk.Label(nav,text='CHỌN CHỨC NĂNG',font=('Segoe UI',10,'bold')).pack(anchor='w',pady=(3,12))
        self.nav=tk.Listbox(nav,exportselection=False,activestyle='none',borderwidth=0,highlightthickness=0,
                            bg='#f3f5f7',fg='#253f4d',selectbackground='#d1e7e3',selectforeground='#104b45',
                            font=('Segoe UI',11),width=24,height=15)
        for name in TOOLS:self.nav.insert('end',name)
        self.nav.pack(fill='both',expand=True)
        self.nav.bind('<<ListboxSelect>>',lambda e:self.select_tool(self.nav.get(self.nav.curselection()[0])) if self.nav.curselection() else None)
        ttk.Label(nav,text='Nhập số thập phân bằng dấu chấm.\nCtrl+Enter: tính\nCtrl+S: lưu bài',foreground='#455c69',wraplength=200).pack(anchor='w',pady=16)
        input_panel=ttk.Frame(main,width=390); main.add(input_panel,weight=0)
        self.title_label=ttk.Label(input_panel,text='',font=('Segoe UI',16,'bold')); self.title_label.pack(anchor='w',pady=(0,10))
        self.method=tk.StringVar()
        self.combo=ttk.Combobox(input_panel,textvariable=self.method,state='readonly',width=34)
        self.combo.pack(fill='x',pady=(0,8)); self.combo.bind('<<ComboboxSelected>>',lambda e:self.show_form())
        scrollframe=ttk.Frame(input_panel); scrollframe.pack(fill='both',expand=True)
        self.form_canvas=tk.Canvas(scrollframe,bg='#f3f5f7',highlightthickness=0,width=370)
        scroll=ttk.Scrollbar(scrollframe,orient='vertical',command=self.form_canvas.yview)
        self.form_canvas.configure(yscrollcommand=scroll.set); scroll.pack(side='right',fill='y'); self.form_canvas.pack(side='left',fill='both',expand=True)
        self.form_body=ttk.Frame(self.form_canvas); self.form_window=self.form_canvas.create_window((0,0),window=self.form_body,anchor='nw')
        self.form_body.bind('<Configure>',lambda e:self.form_canvas.configure(scrollregion=self.form_canvas.bbox('all')))
        self.form_canvas.bind('<Configure>',lambda e:self.form_canvas.itemconfigure(self.form_window,width=e.width))
        self.bind_all('<MouseWheel>',self.scroll_form)
        self.run_button=ttk.Button(input_panel,text='Tính kết quả   Ctrl+Enter',style='Accent.TButton',command=self.run)
        self.run_button.pack(fill='x',pady=(12,4))
        ttk.Button(input_panel,text='Nạp lại ví dụ của phương pháp',command=self.reset_example).pack(fill='x')
        result=ttk.Frame(main); main.add(result,weight=1)
        toolbar=ttk.Frame(result); toolbar.pack(fill='x',pady=(0,9))
        ttk.Label(toolbar,text='Kết quả',font=('Segoe UI',16,'bold')).pack(side='left')
        ttk.Button(toolbar,text='Xuất bảng',command=self.export_tables).pack(side='right',padx=(5,0))
        ttk.Button(toolbar,text='Lưu lời giải',command=self.export_report).pack(side='right',padx=5)
        ttk.Button(toolbar,text='Sao chép',command=self.copy_report).pack(side='right')
        precision_row=ttk.Frame(result);precision_row.pack(fill='x',pady=(0,8))
        ttk.Label(precision_row,text='Số chữ số thập phân (0–15):').pack(side='left')
        self.digits=tk.StringVar(value='7')
        ttk.Spinbox(precision_row,from_=0,to=15,textvariable=self.digits,width=4).pack(side='left',padx=8)
        ttk.Label(precision_row,text='Áp dụng khi bấm Tính kết quả; không làm tròn trong phép tính.').pack(side='left')
        self.notebook=ttk.Notebook(result); self.notebook.pack(fill='both',expand=True)
        self.math_document=MathDocument(self.notebook)
        self.notebook.add(self.math_document,text='Lời giải toán học')
        self.output=ScrolledText(self.notebook,font=('Consolas',11),wrap='none',bg='white',fg='#172b3a',padx=16,pady=16,undo=False)
        self.notebook.add(self.output,text='Văn bản sao chép')
        table_page=ttk.Frame(self.notebook); self.notebook.add(table_page,text='Bảng số')
        self.table_select=ttk.Combobox(table_page,state='readonly'); self.table_select.pack(fill='x',padx=8,pady=8)
        self.table_select.bind('<<ComboboxSelected>>',lambda e:self.show_table())
        modes=ttk.Frame(table_page);modes.pack(fill='x',padx=8)
        self.table_mode=tk.StringVar(value='math')
        ttk.Radiobutton(modes,text='Trình bày theo tài liệu',variable=self.table_mode,value='math',command=self.show_table).pack(side='left')
        ttk.Radiobutton(modes,text='Dữ liệu dạng lưới',variable=self.table_mode,value='data',command=self.show_table).pack(side='left',padx=12)
        self.table_note=ttk.Label(table_page,text='',wraplength=700);self.table_note.pack(fill='x',padx=8,pady=5)
        self.math_table=MathTable(table_page)
        table_frame=ttk.Frame(table_page); table_frame.pack(fill='both',expand=True)
        self.table_frame=table_frame
        self.tree=ttk.Treeview(table_frame,show='headings'); self.tree.grid(row=0,column=0,sticky='nsew')
        sv=ttk.Scrollbar(table_frame,orient='vertical',command=self.tree.yview); sv.grid(row=0,column=1,sticky='ns')
        sh=ttk.Scrollbar(table_frame,orient='horizontal',command=self.tree.xview); sh.grid(row=1,column=0,sticky='ew')
        self.tree.configure(yscrollcommand=sv.set,xscrollcommand=sh.set); table_frame.rowconfigure(0,weight=1); table_frame.columnconfigure(0,weight=1)
        plot_page=ttk.Frame(self.notebook); self.notebook.add(plot_page,text='Đồ thị')
        self.fig=Figure(figsize=(6,4),dpi=100,facecolor='white'); self.ax=self.fig.add_subplot(111)
        self.plot=FigureCanvasTkAgg(self.fig,master=plot_page); self.plot.get_tk_widget().pack(fill='both',expand=True)
        NavigationToolbar2Tk(self.plot,plot_page).update()
        self.status=tk.StringVar(value='Sẵn sàng. Dữ liệu đang hiển thị là ví dụ; bạn có thể thay toàn bộ.')
        ttk.Label(self,textvariable=self.status,padding=(18,8),foreground='#344e5a').pack(fill='x')
        self.set_output('Nhập dữ liệu ở bên trái rồi chọn Tính kết quả.\n\nBạn có thể gõ trực tiếp, dán các dãy số hoặc nhập Excel/CSV.\nCác kết quả có thể xuất ra Excel, CSV và tệp lời giải.')

    def scroll_form(self,event):
        widget=self.winfo_containing(event.x_root,event.y_root)
        while widget:
            if widget==self.form_body or widget==self.form_canvas:
                self.form_canvas.yview_scroll(-int(event.delta/120),'units'); return
            widget=getattr(widget,'master',None)

    def snapshot(self):
        return {'tool':self.current_tool,'method':self.current_method,
                'values':{**{k:v.get() for k,v in self.form.items()},'digits':self.digits.get()},
                'x':self.xinput.get('1.0','end').strip() if hasattr(self,'xinput') and self.xinput.winfo_exists() else '',
                'y':self.yinput.get('1.0','end').strip() if hasattr(self,'yinput') and self.yinput.winfo_exists() else ''}

    def remember(self):
        if self.current_tool and self.current_method: self.cache[(self.current_tool,self.current_method)]=self.snapshot()

    def select_tool(self,name):
        if self.current_tool==name:return
        self.remember(); self.current_tool=name; self.current_method=None
        self.combo.configure(values=TOOLS[name][0]); self.method.set(TOOLS[name][0][0]); self.title_label.configure(text=name)
        i=list(TOOLS).index(name); self.nav.selection_clear(0,'end'); self.nav.selection_set(i)
        self.show_form(save=False)

    def show_form(self,save=True):
        previous=self.snapshot() if save else {}
        if save:self.remember()
        self.current_method=self.method.get()
        saved=self.cache.get((self.current_tool,self.current_method),{})
        self.view_revision=getattr(self,'view_revision',0)+1
        self.report=None
        self.table_select.configure(values=[]);self.table_select.set('')
        self.tree.delete(*self.tree.get_children());self.math_table.canvas.delete('all')
        self.ax.clear();self.plot.draw_idle()
        self.set_output('Đã chọn '+self.current_method+'. Nhấn Tính kết quả để tạo lời giải mới.')
        self.digits.set(saved.get('values',{}).get('digits',self.digits.get()))
        for w in self.form_body.winfo_children():w.destroy()
        self.form={}
        for key,label,default,hint in fields(self.current_tool,self.current_method):
            ttk.Label(self.form_body,text=label,wraplength=355).pack(anchor='w',pady=(9,3))
            var=tk.StringVar(value=saved.get('values',{}).get(key,previous.get('values',{}).get(key,default))); self.form[key]=var
            if key=='direction':
                ttk.Combobox(self.form_body,textvariable=var,values=['Tiến','Lùi'],state='readonly').pack(fill='x',padx=(0,8))
            else: ttk.Entry(self.form_body,textvariable=var).pack(fill='x',padx=(0,8))
            if hint:ttk.Label(self.form_body,text=hint,foreground='#4c626e',wraplength=350,font=('Segoe UI',9)).pack(anchor='w',pady=(2,0))
        needs_data=TOOLS[self.current_tool][1]
        if self.current_method in ('Chebyshev','Horner'):needs_data=False
        if self.current_tool=='Tích phân' and self.current_method!='Tích phân kép':needs_data=True
        if needs_data:
            ttk.Separator(self.form_body).pack(fill='x',pady=(18,8))
            ttk.Label(self.form_body,text='BẢNG DỮ LIỆU',font=('Segoe UI',10,'bold')).pack(anchor='w')
            ttk.Label(self.form_body,text='Mỗi dãy gồm các số cách nhau bằng dấu cách.\nMốc trùng: giữ điểm xuất hiện đầu tiên.',wraplength=350,foreground='#4c626e').pack(anchor='w',pady=5)
            for key,label,default in [('x','Dãy x','0 1 2 3 4'),('y','Dãy y','0 1 4 9 16')]:
                ttk.Label(self.form_body,text=label).pack(anchor='w',pady=(7,3))
                widget=tk.Text(self.form_body,height=3,wrap='word',font=('Consolas',11),bg='white',fg='#172b3a',relief='solid',borderwidth=1,undo=True)
                widget.insert('1.0',saved.get(key,previous.get(key) or default)); widget.pack(fill='x',padx=(0,8))
                setattr(self,key+'input',widget)
            ttk.Button(self.form_body,text='Nhập Excel / CSV',command=self.import_data).pack(anchor='w',pady=9)
            if self.current_method=='Bessel' and not saved and not previous.get('x'):
                self.xinput.delete('1.0','end'); self.xinput.insert('1.0','0 1 2 3')
                self.yinput.delete('1.0','end'); self.yinput.insert('1.0','0 1 4 9')
            if self.current_tool=='Nội suy ngược' and not saved and not previous.get('x'):
                self.xinput.delete('1.0','end'); self.xinput.insert('1.0','1 1.1 1.2 1.3 1.4')
                self.yinput.delete('1.0','end'); self.yinput.insert('1.0','1 1.21 1.44 1.69 1.96')
        self.form_canvas.yview_moveto(0)

    def reset_example(self):
        self.cache.pop((self.current_tool,self.current_method),None); self.show_form(save=False)
        self.status.set('Đã nạp ví dụ. Bạn có thể sửa tất cả ô nhập trước khi tính.')

    def run(self):
        if self.future and not self.future.done():return
        try: set_display_digits(self.digits.get())
        except ValueError:
            self.status.set('Nhập số chữ số thập phân nguyên từ 0 đến 15.');return
        self.run_revision=self.view_revision
        state=self.snapshot(); self.remember()
        self.run_button.configure(state='disabled'); self.status.set('Đang tính '+state['method']+'…')
        self.future=self.executor.submit(dispatch,state['tool'],state['method'],state['values'],state['x'],state['y'])
        self.after(80,self.poll)

    def poll(self):
        if not self.future.done():self.after(100,self.poll);return
        self.run_button.configure(state='normal')
        if self.run_revision!=self.view_revision:
            self.status.set('Phương pháp đã đổi; nhấn Tính kết quả để tính lại.');return
        try:self.display(self.future.result())
        except Exception as exc:
            self.status.set('Chưa tính được — kiểm tra dữ liệu nhập.')
            self.report=None; self.table_select.configure(values=[]); self.table_select.set(''); self.tree.delete(*self.tree.get_children())
            self.math_table.canvas.delete('all');self.table_note.configure(text='')
            self.ax.clear(); self.plot.draw_idle()
            self.set_output('Không tính được kết quả\n\n'+str(exc)+'\n\nSửa dữ liệu rồi nhấn Tính kết quả.'); self.notebook.select(0)
            with (ROOT/'error.log').open('a',encoding='utf-8') as out:out.write(traceback.format_exc()+'\n')

    def display(self,report):
        self.report=report; self.set_output(report.text()); self.notebook.select(0)
        self.math_document.show(report)
        self.table_select.configure(values=[v[0] for v in report.tables])
        if report.tables:self.table_select.current(0);self.show_table()
        else:
            self.table_select.set('');self.tree.delete(*self.tree.get_children());self.math_table.canvas.delete('all')
        self.ax.clear()
        for xx,yy,label in report.curves:self.ax.plot(xx,yy,label=label,linewidth=1.8)
        if report.scatter is not None:self.ax.scatter(*report.scatter,color='#b85b32',s=30,zorder=5,label='Dữ liệu')
        self.ax.set_xlabel('t' if report.title.startswith('Bài toán Cauchy') else 'x')
        self.ax.set_ylabel('Giá trị'); self.ax.grid(True,alpha=.2)
        if report.curves or report.scatter is not None:self.ax.legend(fontsize=9)
        else:self.ax.text(.5,.5,'Phương pháp này không có đồ thị.',ha='center',transform=self.ax.transAxes)
        self.fig.tight_layout();self.plot.draw_idle()
        self.status.set('Hoàn tất: '+report.title+f' · Hiển thị tối đa {self.digits.get()} chữ số thập phân; tính bằng độ chính xác đầy đủ.')

    def show_table(self):
        if not self.report or not self.report.tables:return
        i=self.table_select.current()
        if i<0:return
        _,headers,rows=self.report.tables[i]; self.tree.delete(*self.tree.get_children()); cols=[str(k) for k in range(len(headers))]
        if self.table_mode.get()=='math':
            self.table_frame.pack_forget();self.math_table.pack(fill='both',expand=True)
            self.math_table.show(self.report,i)
            self.table_note.configure(text='Sai phân cùng hàng với mốc bắt đầu, chụm lên trên; ô viền xanh là giá trị được lấy vào công thức. Cuộn ngang để xem các cấp bên phải.' if i==0 and self.report.presentation.get('difference') else 'Công thức và chỉ số được dựng từ LaTeX. Bảng hệ số: bậc cao bên trái, hệ số tự do bên phải.')
            if len(rows)>80 or len(headers)>12 or (i==0 and self.report.presentation.get('difference') and len(rows)>11):
                self.table_note.configure(text=self.table_note.cget('text')+' Bản trình bày rút gọn; chọn Dữ liệu dạng lưới hoặc xuất Excel để xem đầy đủ.')
        else:
            self.math_table.pack_forget();self.table_frame.pack(fill='both',expand=True)
            self.table_note.configure(text='Dữ liệu số gốc để tra cứu; chọn “Trình bày theo tài liệu” để xem bố cục toán học.')
        self.tree.configure(columns=cols)
        for col,label in zip(cols,headers):self.tree.heading(col,text=label);self.tree.column(col,width=135,minwidth=90,anchor='e')
        for row in rows[:3000]:self.tree.insert('','end',values=[fmt(v) for v in row])

    def set_output(self,text):
        self.math_document.message(text)
        self.output.configure(state='normal');self.output.delete('1.0','end');self.output.insert('1.0',text);self.output.configure(state='disabled')

    def import_data(self):
        path=filedialog.askopenfilename(filetypes=[('Bảng dữ liệu','*.xlsx *.csv *.txt'),('Tất cả','*.*')])
        if not path:return
        try:
            if Path(path).suffix.lower()=='.xlsx':
                book=load_workbook(path,read_only=True,data_only=True)
                if len(book.sheetnames)>1:
                    from tkinter.simpledialog import askstring
                    sheet=askstring('Chọn sheet','Các sheet: '+', '.join(book.sheetnames),initialvalue=book.sheetnames[0],parent=self)
                    if not sheet:book.close();return
                    if sheet not in book.sheetnames:book.close();raise ValueError('Tên sheet không tồn tại.')
                else:sheet=book.sheetnames[0]
                rows=list(book[sheet].iter_rows(values_only=True));book.close()
            else:
                text=Path(path).read_text(encoding='utf-8-sig')
                try:dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t');rows=list(csv.reader(text.splitlines(),dialect))
                except csv.Error:rows=[line.split() for line in text.splitlines()]
            data=[]
            for i,row in enumerate(rows,1):
                if not row or all(v is None or str(v).strip()=='' for v in row):continue
                try:
                    pair=(float(row[0]),float(row[1]))
                    if not np.all(np.isfinite(pair)):raise ValueError()
                    data.append(pair)
                except (ValueError,TypeError,IndexError):
                    if not data and i==1:continue
                    raise ValueError(f'Dòng {i} không có đủ hai số x,y hữu hạn ở hai cột đầu.')
            if not data:raise ValueError('Không có dữ liệu. Hai cột đầu phải là x và y.')
            for i,w in enumerate((self.xinput,self.yinput)):
                w.delete('1.0','end');w.insert('1.0',' '.join(repr(v[i]) for v in data))
            self.status.set(f'Đã nhập {len(data)} điểm từ {Path(path).name}.')
        except Exception as exc:messagebox.showerror('Không nhập được dữ liệu',str(exc),parent=self)

    def save_project(self):
        path=filedialog.asksaveasfilename(defaultextension='.json',filetypes=[('Bài tính PPS','*.json')])
        if path:
            try:Path(path).write_text(json.dumps(self.snapshot(),ensure_ascii=False,indent=2),encoding='utf-8');self.status.set('Đã lưu bài tính: '+path)
            except OSError as exc:messagebox.showerror('Không lưu được',str(exc),parent=self)

    def load_project(self):
        path=filedialog.askopenfilename(filetypes=[('Bài tính PPS','*.json')])
        if not path:return
        try:
            state=json.loads(Path(path).read_text(encoding='utf-8'))
            if state['tool'] not in TOOLS or state['method'] not in TOOLS[state['tool']][0] or not isinstance(state['values'],dict):raise ValueError('Tệp không đúng cấu trúc bài tính PPS.')
            self.remember();self.cache[(state['tool'],state['method'])]=state
            # Bypass remembering an older form over the imported state.
            self.current_tool=state['tool']; self.current_method=state['method']
            self.combo.configure(values=TOOLS[self.current_tool][0]);self.method.set(self.current_method);self.title_label.configure(text=self.current_tool)
            self.nav.selection_clear(0,'end');self.nav.selection_set(list(TOOLS).index(self.current_tool));self.show_form(save=False)
            self.status.set('Đã mở bài tính. Nhấn Tính kết quả để chạy.')
        except Exception as exc:messagebox.showerror('Không mở được bài',str(exc),parent=self)

    def copy_report(self):
        if self.report:self.clipboard_clear();self.clipboard_append(self.report.text());self.status.set('Đã sao chép lời giải.')

    def export_report(self):
        if not self.report:self.status.set('Hãy tính kết quả trước khi xuất.');return
        path=filedialog.asksaveasfilename(defaultextension='.html',filetypes=[('Lời giải toán học (mở bằng trình duyệt)','*.html'),('Văn bản thuần','*.txt')])
        if path:
            try:
                if Path(path).suffix.lower()=='.html':export_html(self.report,Path(path))
                else:Path(path).write_text(self.report.text(),encoding='utf-8')
                self.status.set('Đã lưu lời giải.')
            except OSError as exc:messagebox.showerror('Không lưu được',str(exc),parent=self)

    def export_tables(self):
        if not self.report or not self.report.tables:self.status.set('Chưa có bảng để xuất.');return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',filetypes=[('Toàn bộ bảng Excel','*.xlsx'),('Bảng đang chọn CSV','*.csv')])
        if not path:return
        try:
            def clean(v):
                if v is None:return None
                if isinstance(v,str):return "'"+v if v.startswith(('=','+','-','@')) else v
                return float(v)
            if Path(path).suffix.lower()=='.csv':
                index=max(0,self.table_select.current());_,headers,rows=self.report.tables[index]
                with open(path,'w',encoding='utf-8-sig',newline='') as out:
                    writer=csv.writer(out);writer.writerow(headers);writer.writerows([[clean(v) for v in row] for row in rows])
            else:
                book=Workbook();book.remove(book.active)
                for i,(name,headers,rows) in enumerate(self.report.tables,1):
                    title=''.join(c for c in name if c not in '[]:*?/\\')
                    sheet=book.create_sheet((str(i)+' '+title)[:31]);sheet.append(headers)
                    for row in rows:sheet.append([clean(v) for v in row])
                    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
                    for column in sheet.columns:
                        sheet.column_dimensions[column[0].column_letter].width=20
                        for cell in column[1:]:
                            if cell.data_type=='n':cell.number_format='0.#######'
                book.save(path)
            self.status.set('Đã xuất toàn bộ dữ liệu, giữ độ chính xác số: '+path)
        except Exception as exc:messagebox.showerror('Không xuất được',str(exc),parent=self)

    def help(self):
        win=tk.Toplevel(self);win.title('Hướng dẫn PPS');win.geometry('800x650')
        text=ScrolledText(win,wrap='word',font=('Segoe UI',11),padx=20,pady=20);text.pack(fill='both',expand=True)
        text.insert('1.0',(ROOT/'HUONG_DAN.md').read_text(encoding='utf-8'));text.configure(state='disabled')

    def close(self):
        self.executor.shutdown(wait=False,cancel_futures=True);self.destroy()


def main():
    App().mainloop()
