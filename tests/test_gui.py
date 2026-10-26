"""Desktop integration tests; Tk window stays withdrawn during tests."""
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from pps.gui import App, TOOLS, dispatch


class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=App();cls.app.withdraw()

    @classmethod
    def tearDownClass(cls):cls.app.close()

    def test_every_method_accepts_its_editable_example(self):
        for tool,(methods,_) in TOOLS.items():
            self.app.select_tool(tool)
            for method in methods:
                with self.subTest(tool=tool,method=method):
                    self.app.method.set(method);self.app.show_form()
                    self.app.reset_example()
                    s=self.app.snapshot()
                    r=dispatch(tool,method,s['values'],s['x'],s['y'])
                    self.assertTrue(r.title)
                    self.app.display(r)
                    self.app.update_idletasks()

    def test_excel_csv_import_export_and_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp=Path(tmp); source=tmp/'input.xlsx'
            wb=Workbook();ws=wb.active;ws.append(['x','y'])
            for x in range(3):ws.append([x,x*x])
            wb.save(source)
            self.app.select_tool('Nội suy đa thức');self.app.method.set('Newton');self.app.show_form()
            self.app.form['newton_nodes'].set('Mốc bất kỳ')
            with patch('pps.gui.filedialog.askopenfilename',return_value=str(source)):
                self.app.import_data()
            s=self.app.snapshot();self.assertEqual(s['x'],'0.0 1.0 2.0')
            self.app.display(dispatch(s['tool'],s['method'],s['values'],s['x'],s['y']))
            out=tmp/'output.xlsx'
            with patch('pps.gui.filedialog.asksaveasfilename',return_value=str(out)):self.app.export_tables()
            wb=load_workbook(out);self.assertEqual(len(wb.sheetnames),2);self.assertEqual(wb.worksheets[0]['A2'].value,0);wb.close()
            out=tmp/'output.csv'
            with patch('pps.gui.filedialog.asksaveasfilename',return_value=str(out)):self.app.export_tables()
            self.assertIn('0.0',out.read_text(encoding='utf-8-sig'))
            project=tmp/'project.json'
            with patch('pps.gui.filedialog.asksaveasfilename',return_value=str(project)):self.app.save_project()
            self.app.select_tool('Bài toán biên')
            with patch('pps.gui.filedialog.askopenfilename',return_value=str(project)):self.app.load_project()
            self.assertEqual(self.app.snapshot(),json.loads(project.read_text(encoding='utf-8')))

    def test_newton_variants_precision_and_switch(self):
        from pps.common import set_display_digits,fmt
        import numpy as np
        for nodes in ('Mốc cách đều','Mốc bất kỳ'):
            for direction in ('Tiến','Lùi'):
                r=dispatch('Nội suy đa thức','Newton',{'target':'1.5','order':'0','direction':direction,'newton_nodes':nodes},'0 1 2 3','1 2 5 10')
                self.assertAlmostEqual(r.value(1.5),3.25)
        self.assertNotIn('Newton',TOOLS)
        r=dispatch('Newton','Mốc bất kỳ',{'target':'1.5','order':'0','direction':'Lùi'},'0 0.7 2 3','1 1.49 5 10')
        self.assertAlmostEqual(r.value(1.5),3.25)
        set_display_digits(0);self.assertEqual(fmt(100), '100')
        set_display_digits(10);self.assertEqual(fmt(1/3),'0.3333333333')
        set_display_digits(7)
        self.app.select_tool('Nội suy đa thức')
        self.app.display(r)
        self.app.method.set('Gauss I');self.app.show_form()
        self.assertIsNone(self.app.report)
        self.assertEqual(len(self.app.tree.get_children()),0)

    def test_interpolation_builds_polynomial_without_a_target(self):
        report=dispatch('Nội suy đa thức','Lagrange',{'target':'','order':'0'},
                        '1.2 1.5 1.7 1.8 2.1 2.3',
                        '.892 1.179 1.358 1.445 1.688 1.839')
        self.assertAlmostEqual(report.value(0), -34011/5500, places=13)

    def test_async_calculate_button(self):
        self.app.select_tool('Nội suy đa thức');self.app.reset_example();self.app.run()
        deadline=time.monotonic()+20
        while self.app.future and not self.app.future.done() and time.monotonic()<deadline:
            self.app.update();time.sleep(.02)
        self.assertTrue(self.app.future.done())
        self.app.poll();self.assertIsNotNone(self.app.report)


if __name__=='__main__':unittest.main()
